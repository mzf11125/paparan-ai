"""PDF and PPTX export tools for policy briefs."""
import io
from app.db.schema import PolicyBrief

WATERMARK_COLORS = {
    "unclassified": (0.85, 0.85, 0.85),
    "official": (0.7, 0.8, 1.0),
    "confidential": (1.0, 0.85, 0.6),
    "secret": (1.0, 0.7, 0.7),
}


def _draw_watermark(canvas, text: str, color: tuple):
    """Draw diagonal watermark on a ReportLab canvas page."""
    from reportlab.lib.pagesizes import A4
    canvas.saveState()
    canvas.setFillColorRGB(*color)
    canvas.setFont("Helvetica-Bold", 60)
    canvas.translate(A4[0] / 2, A4[1] / 2)
    canvas.rotate(45)
    canvas.setFillAlpha(0.15)
    canvas.drawCentredString(0, 0, text.upper())
    canvas.restoreState()


def generate_pdf(brief: PolicyBrief, watermark: str = None) -> bytes:
    """Generate a diplomat-ready PDF from a PolicyBrief.

    Args:
        brief: PolicyBrief instance
        watermark: Override watermark text (defaults to classification level)

    Returns:
        PDF bytes
    """
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.enums import TA_CENTER

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=2*cm, bottomMargin=2*cm,
                            leftMargin=2.5*cm, rightMargin=2.5*cm)
    styles = getSampleStyleSheet()
    wm_text = watermark or brief.classification.upper()
    wm_color = WATERMARK_COLORS.get(brief.classification, (0.85, 0.85, 0.85))

    def on_page(canvas, doc):
        _draw_watermark(canvas, wm_text, wm_color)
        # Classification banner top
        canvas.setFillColorRGB(0.1, 0.2, 0.5)
        canvas.rect(0, A4[1] - 1*cm, A4[0], 1*cm, fill=1, stroke=0)
        canvas.setFillColorRGB(1, 1, 1)
        canvas.setFont("Helvetica-Bold", 9)
        canvas.drawCentredString(A4[0]/2, A4[1] - 0.65*cm, brief.classification.upper())
        # Classification banner bottom
        canvas.setFillColorRGB(0.1, 0.2, 0.5)
        canvas.rect(0, 0, A4[0], 0.8*cm, fill=1, stroke=0)
        canvas.setFillColorRGB(1, 1, 1)
        canvas.drawCentredString(A4[0]/2, 0.25*cm, brief.classification.upper())

    h1 = ParagraphStyle("h1", parent=styles["Heading1"], fontSize=16, spaceAfter=6)
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], fontSize=11, spaceAfter=4)
    body = ParagraphStyle("body", parent=styles["Normal"], fontSize=9, spaceAfter=4, leading=14)

    story = [
        Paragraph(f"POLICY INTELLIGENCE BRIEF", ParagraphStyle("label", parent=styles["Normal"],
                  fontSize=8, textColor=colors.grey, spaceAfter=2)),
        Paragraph(brief.title, h1),
        Paragraph(f"{brief.region} · {brief.date} · {brief.classification.upper()}", body),
        Spacer(1, 0.4*cm),
        Paragraph("EXECUTIVE SUMMARY", h2),
    ]
    for i, item in enumerate(brief.executiveSummary, 1):
        story.append(Paragraph(f"{i}. {item}", body))

    story += [Spacer(1, 0.3*cm), Paragraph("CURRENT SITUATION", h2),
              Paragraph(brief.currentSituation, body),
              Spacer(1, 0.3*cm), Paragraph("STRATEGIC IMPLICATIONS", h2),
              Paragraph(brief.implications, body),
              Spacer(1, 0.3*cm), Paragraph("RECOMMENDED ACTIONS", h2)]

    high_actions = [a for a in brief.actions if a.priority == "HIGH"] or brief.actions[:3]
    tdata = [["Priority", "Action", "Owner", "Deadline"]] + \
            [[a.priority, a.text[:80], a.owner or "—", a.deadline or "—"] for a in high_actions]
    t = Table(tdata, colWidths=[2*cm, 9*cm, 3*cm, 3*cm])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1a3a6b")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.lightgrey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f7fa")]),
    ]))
    story.append(t)

    story += [Spacer(1, 0.3*cm), Paragraph("SOURCES", h2)]
    for s in brief.sources:
        story.append(Paragraph(f"• {s.title} [{s.confidence}] — {s.url}", body))

    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)
    return buf.getvalue()


def generate_pptx(brief: PolicyBrief) -> bytes:
    """Generate a 5-slide PPTX from a PolicyBrief.

    Returns:
        PPTX bytes
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN

    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]  # blank layout

    NAVY = RGBColor(0x1a, 0x3a, 0x6b)
    WHITE = RGBColor(0xFF, 0xFF, 0xFF)
    GREY = RGBColor(0x66, 0x66, 0x66)

    def add_text_box(slide, text, left, top, width, height, size=18, bold=False, color=None, align=PP_ALIGN.LEFT):
        txb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        tf = txb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = align
        run = p.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = bold
        if color:
            run.font.color.rgb = color
        return txb

    def add_header_bar(slide, title):
        bar = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(1.2))
        bar.fill.solid()
        bar.fill.fore_color.rgb = NAVY
        bar.line.fill.background()
        tf = bar.text_frame
        tf.text = title
        tf.paragraphs[0].runs[0].font.color.rgb = WHITE
        tf.paragraphs[0].runs[0].font.size = Pt(24)
        tf.paragraphs[0].runs[0].font.bold = True

    def add_watermark(slide, text):
        txb = slide.shapes.add_textbox(Inches(2), Inches(2.5), Inches(9), Inches(3))
        tf = txb.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = text.upper()
        run.font.size = Pt(72)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0xDD, 0xDD, 0xDD)
        txb.rotation = -45

    # Slide 1: Title
    s1 = prs.slides.add_slide(blank)
    add_watermark(s1, brief.classification)
    bar = s1.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
    bar.fill.solid(); bar.fill.fore_color.rgb = NAVY; bar.line.fill.background()
    add_text_box(s1, "POLICY INTELLIGENCE BRIEF", 0.5, 1.5, 12, 0.6, 12, color=RGBColor(0xAA, 0xBB, 0xDD), align=PP_ALIGN.CENTER)
    add_text_box(s1, brief.title, 0.5, 2.2, 12, 2, 28, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_text_box(s1, f"{brief.region}  ·  {brief.date}  ·  {brief.classification.upper()}", 0.5, 4.5, 12, 0.5, 14, color=RGBColor(0xAA, 0xBB, 0xDD), align=PP_ALIGN.CENTER)

    # Slide 2: Executive Summary
    s2 = prs.slides.add_slide(blank)
    add_watermark(s2, brief.classification)
    add_header_bar(s2, "Executive Summary")
    for i, item in enumerate(brief.executiveSummary[:5]):
        add_text_box(s2, f"• {item}", 0.5, 1.4 + i * 1.0, 12, 0.9, 14, color=RGBColor(0x22, 0x22, 0x22))

    # Slide 3: Current Situation
    s3 = prs.slides.add_slide(blank)
    add_watermark(s3, brief.classification)
    add_header_bar(s3, "Current Situation")
    add_text_box(s3, brief.currentSituation[:600], 0.5, 1.4, 12, 5, 13, color=RGBColor(0x22, 0x22, 0x22))

    # Slide 4: Key Developments & Actions
    s4 = prs.slides.add_slide(blank)
    add_watermark(s4, brief.classification)
    add_header_bar(s4, "Key Developments & Recommended Actions")
    for i, d in enumerate(brief.developments[:3]):
        add_text_box(s4, f"[{d.impact}] {d.text[:120]}", 0.5, 1.4 + i * 0.8, 12, 0.7, 12)
    add_text_box(s4, "ACTIONS", 0.5, 4.0, 12, 0.4, 11, bold=True, color=NAVY)
    for i, a in enumerate([x for x in brief.actions if x.priority == "HIGH"][:3]):
        add_text_box(s4, f"→ {a.text[:120]}", 0.5, 4.5 + i * 0.6, 12, 0.5, 11)

    # Slide 5: Sources
    s5 = prs.slides.add_slide(blank)
    add_watermark(s5, brief.classification)
    add_header_bar(s5, "Sources & Confidence")
    for i, src in enumerate(brief.sources[:8]):
        add_text_box(s5, f"[{src.confidence}] {src.title}", 0.5, 1.4 + i * 0.65, 12, 0.6, 11)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def generate_diplomat_pdf(brief: PolicyBrief, to: str, from_name: str,
                          ref_no: str, distribution: list[str] = None) -> bytes:
    """Generate a formal diplomatic memo PDF.

    Args:
        brief: PolicyBrief instance
        to: Recipient name/title
        from_name: Sender name/title
        ref_no: Reference number e.g. PAP-2026-0041
        distribution: List of distribution recipients

    Returns:
        PDF bytes
    """
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable, Table, TableStyle

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=2.5*cm, bottomMargin=2.5*cm,
                            leftMargin=3*cm, rightMargin=3*cm)
    styles = getSampleStyleSheet()
    wm_color = WATERMARK_COLORS.get(brief.classification, (0.85, 0.85, 0.85))

    def on_page(canvas, doc):
        _draw_watermark(canvas, brief.classification.upper(), wm_color)
        # Top banner
        canvas.setFillColorRGB(0.1, 0.2, 0.5)
        canvas.rect(0, A4[1] - 1.2*cm, A4[0], 1.2*cm, fill=1, stroke=0)
        canvas.setFillColorRGB(1, 1, 1)
        canvas.setFont("Helvetica-Bold", 10)
        canvas.drawCentredString(A4[0]/2, A4[1] - 0.75*cm,
                                 f"{'—' * 8}  {brief.classification.upper()}  {'—' * 8}")
        # Bottom banner
        canvas.setFillColorRGB(0.1, 0.2, 0.5)
        canvas.rect(0, 0, A4[0], 1*cm, fill=1, stroke=0)
        canvas.setFillColorRGB(1, 1, 1)
        canvas.setFont("Helvetica-Bold", 9)
        canvas.drawCentredString(A4[0]/2, 0.35*cm,
                                 f"{brief.classification.upper()}  ·  {ref_no}  ·  PAPARAN POLICY INTELLIGENCE")

    label = ParagraphStyle("label", parent=styles["Normal"], fontSize=8,
                           textColor=colors.HexColor("#666666"), spaceAfter=1)
    value = ParagraphStyle("value", parent=styles["Normal"], fontSize=10, spaceAfter=6, leading=14)
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], fontSize=10, spaceAfter=4,
                        textColor=colors.HexColor("#1a3a6b"))
    body = ParagraphStyle("body", parent=styles["Normal"], fontSize=9, spaceAfter=4, leading=14)

    # Memo header block
    header_data = [
        [Paragraph("TO:", label), Paragraph(to, value)],
        [Paragraph("FROM:", label), Paragraph(from_name, value)],
        [Paragraph("SUBJECT:", label), Paragraph(brief.title, value)],
        [Paragraph("DATE:", label), Paragraph(brief.date, value)],
        [Paragraph("REF:", label), Paragraph(ref_no, value)],
        [Paragraph("CLASSIFICATION:", label), Paragraph(brief.classification.upper(), value)],
    ]
    ht = Table(header_data, colWidths=[3.5*cm, 12*cm])
    ht.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LINEBELOW", (0, -1), (-1, -1), 0.5, colors.HexColor("#1a3a6b")),
    ]))

    story = [Spacer(1, 0.3*cm), ht, Spacer(1, 0.5*cm),
             Paragraph("EXECUTIVE SUMMARY", h2)]

    for i, item in enumerate(brief.executiveSummary[:3], 1):
        story.append(Paragraph(f"{i}. {item}", body))

    story += [Spacer(1, 0.3*cm), Paragraph("RECOMMENDED ACTIONS (HIGH PRIORITY)", h2)]
    high_actions = [a for a in brief.actions if a.priority == "HIGH"] or brief.actions[:3]
    for a in high_actions:
        story.append(Paragraph(f"→ {a.text}", body))

    if distribution:
        story += [Spacer(1, 0.4*cm), HRFlowable(width="100%", thickness=0.5),
                  Spacer(1, 0.2*cm),
                  Paragraph("DISTRIBUTION:", label),
                  Paragraph(" · ".join(distribution), body)]

    story += [Spacer(1, 0.5*cm), HRFlowable(width="100%", thickness=0.5),
              Spacer(1, 0.3*cm),
              Paragraph("Prepared by Paparan Policy Intelligence System", label),
              Paragraph(f"Source confidence: {brief.confidence_score} · {brief.source_count} sources verified", label)]

    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)
    return buf.getvalue()
